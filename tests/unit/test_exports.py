"""Geradores de PDF e XLSX (S7, D048): conteúdo, segurança da planilha e texto do PDF."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from telegrana.exports import pdf, xlsx
from telegrana.exports.dados import CategoriaTotal, Dados, Linha

LINHAS = (
    Linha(date(2026, 10, 1), "Mercado do mês", "Mercado", "Pix", "Gasto", 10000, ""),
    Linha(date(2026, 10, 5), '=HYPERLINK("http://x")', "Mercado", "Pix", "Gasto", 5000, ""),
    Linha(date(2026, 10, 5), "Salário", "Salário", "Pix", "Ganho", 300000, ""),
    Linha(date(2026, 10, 6), "tênis", "Vestuário", "Nubank", "Gasto", 20000, "1/3"),
)
DADOS = Dados(
    titulo="Relatório de outubro/2026",
    periodo="outubro/2026",
    visao="Realizado (o que já foi pago)",
    gerado=date(2026, 10, 6),
    entrou=300000,
    saiu=35000,
    categorias=(CategoriaTotal("Vestuário", 20000, 0), CategoriaTotal("Mercado", 15000, 0)),
    lancamentos=LINHAS,
    compromissos=(
        Linha(date(2026, 11, 10), "tênis", "Vestuário", "Nubank", "Gasto", 20000, "2/3"),
    ),
)


def test_planilha_tem_as_quatro_abas_valores_e_neutraliza_formula() -> None:
    wb = load_workbook(io.BytesIO(xlsx.gera(DADOS)))
    assert wb.sheetnames == ["Resumo", "Por categoria", "Lançamentos", "Compromissos"]
    resumo = wb["Resumo"]
    assert resumo["A5"].value == "Entrou"
    assert resumo["B5"].value == Decimal("3000")
    assert resumo["B7"].value == Decimal("2650")  # saldo
    assert resumo["B5"].number_format == xlsx.MOEDA
    lanc = wb["Lançamentos"]
    assert [c.value for c in lanc[1]] == [
        "Data",
        "Descrição",
        "Categoria",
        "Forma",
        "Tipo",
        "Valor",
        "Parcela",
    ]
    assert lanc["B3"].value == '\'=HYPERLINK("http://x")'  # não vira fórmula
    assert lanc["B3"].data_type == "s"
    assert lanc["F5"].value == Decimal("200")
    assert lanc.auto_filter.ref == "A1:G5"
    assert wb["Compromissos"]["G2"].value == "2/3"
    assert xlsx.seguro("-5") == "'-5"
    assert xlsx.seguro("@x") == "'@x"
    assert xlsx.seguro("normal") == "normal"


def test_pdf_e_texto_latin1() -> None:
    conteudo = pdf.gera(DADOS)
    assert conteudo.startswith(b"%PDF")
    assert len(conteudo) > 10_000  # tem a logo
    assert pdf.texto("🛒 Mercado — feira…") == "Mercado - feira..."
    assert pdf.texto("Ação · R$ 10,00") == "Ação · R$ 10,00"


def test_pdf_com_muitos_lancamentos_pagina_e_avisa() -> None:
    muitos = tuple(
        Linha(date(2026, 10, 1 + i % 28), f"compra {i}", "Mercado", "Pix", "Gasto", 100 + i, "")
        for i in range(400)
    )
    conteudo = pdf.gera(_com(muitos))
    assert conteudo.startswith(b"%PDF")


def _com(linhas: tuple[Linha, ...]) -> Dados:
    return Dados(
        titulo=DADOS.titulo,
        periodo=DADOS.periodo,
        visao=DADOS.visao,
        gerado=DADOS.gerado,
        entrou=DADOS.entrou,
        saiu=DADOS.saiu,
        categorias=DADOS.categorias,
        lancamentos=linhas,
        compromissos=(),
        truncado=True,
    )
