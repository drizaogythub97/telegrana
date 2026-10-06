"""Exportação PDF/XLSX (S7, D048) contra Postgres real: conteúdo conferido contra os dados."""

from __future__ import annotations

import io
from dataclasses import replace
from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from telegrana.core import exportacao, relatorios
from telegrana.core import textos as t
from telegrana.core.contexto import FUSO
from telegrana.core.mensagens import Resultado
from telegrana.infra import db
from tests.integration.test_cadastro import CTX, Bot, acoes, conn
from tests.integration.test_lancamentos import IAFalsa, conta
from tests.integration.test_relatorios import com_dados

pytestmark = pytest.mark.integration
__all__ = ["com_dados", "conn"]

HOJE = datetime(2026, 10, 6, 10, tzinfo=FUSO)


@pytest.fixture(autouse=True)
def relogio(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(relatorios, "agora", lambda: HOJE)
    monkeypatch.setattr(exportacao, "agora", lambda: HOJE)


@pytest.fixture
def bot(conn: db.Connection) -> Bot:
    return Bot(conn, replace(CTX, extrator=IAFalsa()))


def planilha(r: Resultado) -> object:
    arquivo = r.saidas[0].arquivo
    assert arquivo is not None
    assert arquivo.nome == "telegrana-2026-10.xlsx"
    return load_workbook(io.BytesIO(arquivo.conteudo))


def test_exportar_por_botoes(bot: Bot, com_dados: str) -> None:
    de = com_dados
    tela = bot(de, comando="exportar")
    assert tela.saidas[0].texto == t.EXPORTAR_TITULO
    assert "ex:ma:pdf" in acoes(tela)
    r = bot(de, acao="ex:ma:pdf")
    arquivo = r.saidas[0].arquivo
    assert arquivo is not None
    assert arquivo.nome == "telegrana-2026-10.pdf"
    assert arquivo.conteudo.startswith(b"%PDF")
    assert r.saidas[0].texto == t.EXPORTACAO_LEGENDA.format(nome="PDF", periodo="outubro/2026")

    wb = planilha(bot(de, acao="ex:ma:xlsx"))
    resumo = wb["Resumo"]  # type: ignore[index]
    assert resumo["B5"].value == Decimal("3000")  # entrou (salário)
    assert resumo["B6"].value == Decimal("180")  # saiu: só o pago (o cartão em aberto fica fora)
    lancamentos = [[c.value for c in linha] for linha in wb["Lançamentos"].iter_rows(min_row=2)]  # type: ignore[index]
    assert len(lancamentos) == 4
    assert {x[1] for x in lancamentos} >= {"feira do sabado", "uber"}
    compromissos = [[c.value for c in linha] for linha in wb["Compromissos"].iter_rows(min_row=2)]  # type: ignore[index]
    assert [(x[5], x[3]) for x in compromissos] == [(Decimal("80"), "Crédito")]
    assert bot(de, acao="ex:xx:pdf").saidas[0].texto == t.USE_OS_BOTOES


def test_exportar_por_frase(bot: Bot, com_dados: str) -> None:
    de = com_dados
    r = bot(de, texto="me manda os gastos de setembro em PDF")
    assert r.rotulo.startswith("exportacao.frase")
    arquivo = r.saidas[0].arquivo
    assert arquivo is not None
    assert arquivo.nome == "telegrana-2026-09.pdf"
    r = bot(de, texto="quero a planilha deste mês")
    wb = planilha(r)
    assert wb["Resumo"]["B5"].value == Decimal("3000")  # type: ignore[index]  # sem "gastos": tudo


def test_fechamento_do_mes_vai_com_pdf(conn: db.Connection, com_dados: str) -> None:
    saidas, falhas = relatorios.da_rotina(conn, "telegram", datetime(2026, 11, 1, 9, tzinfo=FUSO))
    assert falhas == 0
    [mes] = [s for s in saidas if s.destino == com_dados]
    assert mes.texto.startswith(t.MENSAL_TITULO.format(periodo="outubro/2026"))
    assert mes.arquivo is not None
    assert mes.arquivo.nome == "telegrana-2026-10.pdf"


def test_outra_conta_exporta_so_o_que_e_dela(bot: Bot, com_dados: str) -> None:
    outra = conta(bot)
    wb = planilha(bot(outra, acao="ex:ma:xlsx"))
    assert wb["Resumo"]["B6"].value == 0  # type: ignore[index]
    assert list(wb["Lançamentos"].iter_rows(min_row=2)) == []  # type: ignore[index]


def test_arquivo_e_desenhado_fora_da_transacao(
    bot: Bot, conn: db.Connection, com_dados: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Gerar o PDF leva segundos; o papel do app derruba transação parada em 10 s."""
    from psycopg.pq import TransactionStatus

    from telegrana.exports import pdf

    original = pdf.gera
    estados: list[TransactionStatus] = []

    def espia(d: object) -> bytes:
        estados.append(conn.info.transaction_status)
        return original(d)  # type: ignore[arg-type]

    monkeypatch.setattr(pdf, "gera", espia)
    bot(com_dados, acao="ex:ma:pdf")
    relatorios.da_rotina(conn, "telegram", datetime(2026, 11, 1, 9, tzinfo=FUSO))
    assert estados
    assert set(estados) == {TransactionStatus.IDLE}
