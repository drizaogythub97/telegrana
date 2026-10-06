"""Planilha XLSX (PLANO 6; S7, D048): Resumo, Por categoria, Lançamentos e Compromissos.

Moeda BRL, cabeçalho nas cores da marca, filtros e colunas ajustadas. Texto que começa com
`=`, `+`, `-`, `@`, tab ou CR é neutralizado com `'` (injeção de fórmula, PLANO 6).
Gerada em memória: nada vai para disco.
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from telegrana.exports import brand
from telegrana.exports.dados import Dados, Linha

TIPO = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MOEDA = '"R$" #,##0.00'
DATA = "DD/MM/YYYY"
_PERIGOSOS = ("=", "+", "-", "@", "\t", "\r")
_CABECALHO = PatternFill("solid", fgColor=brand.AZUL.lstrip("#"))
_FONTE_CABECALHO = Font(bold=True, color="FFFFFF")


def seguro(texto: str) -> str:
    """Neutraliza injeção de fórmula em XLSX/CSV."""
    return "'" + texto if texto.startswith(_PERIGOSOS) else texto


def _reais(centavos: int) -> Decimal:
    return Decimal(centavos) / 100


def _cabecalho(ws: Worksheet, titulos: Sequence[str], larguras: Sequence[int]) -> None:
    ws.append(list(titulos))
    for i, largura in enumerate(larguras, start=1):
        celula = ws.cell(row=ws.max_row, column=i)
        celula.fill = _CABECALHO
        celula.font = _FONTE_CABECALHO
        celula.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = largura


def _tabela(
    ws: Worksheet,
    titulos: Sequence[str],
    larguras: Sequence[int],
    linhas: Iterable[Sequence[Any]],
    moeda: Sequence[int] = (),
    datas: Sequence[int] = (),
) -> None:
    _cabecalho(ws, titulos, larguras)
    inicio = ws.max_row
    for linha in linhas:
        ws.append([seguro(x) if isinstance(x, str) else x for x in linha])
        for col in moeda:
            ws.cell(row=ws.max_row, column=col).number_format = MOEDA
        for col in datas:
            ws.cell(row=ws.max_row, column=col).number_format = DATA
    ws.freeze_panes = f"A{inicio + 1}"
    if ws.max_row > inicio:
        ws.auto_filter.ref = f"A{inicio}:{get_column_letter(len(titulos))}{ws.max_row}"


def _linhas(lancamentos: Iterable[Linha]) -> Iterable[list[Any]]:
    for x in lancamentos:
        yield [x.dia, x.descricao, x.categoria, x.forma, x.tipo, _reais(x.centavos), x.parcela]


def gera(d: Dados) -> bytes:
    wb = Workbook()
    resumo = wb.active
    if resumo is None:
        raise RuntimeError("planilha sem aba ativa")
    resumo.title = "Resumo"
    resumo.append([seguro(f"Telegrana · {d.titulo}")])
    resumo["A1"].font = Font(bold=True, size=14, color=brand.AZUL.lstrip("#"))
    resumo.append([seguro(d.visao)])
    resumo.append([f"Gerado em {d.gerado:%d/%m/%Y}"])
    resumo.append([])
    for rotulo, valor in (("Entrou", d.entrou), ("Saiu", d.saiu), ("Saldo", d.saldo)):
        resumo.append([rotulo, _reais(valor)])
        resumo.cell(row=resumo.max_row, column=2).number_format = MOEDA
        resumo.cell(row=resumo.max_row, column=1).font = Font(bold=True)
    if d.truncado:
        resumo.append([])
        resumo.append(["Período longo: a aba Lançamentos traz só os mais antigos que couberam."])
    resumo.column_dimensions["A"].width = 30
    resumo.column_dimensions["B"].width = 18

    total = sum(c.gastos for c in d.categorias) or 1
    _tabela(
        wb.create_sheet("Por categoria"),
        ("Categoria", "Gastos", "% dos gastos", "Ganhos"),
        (28, 16, 14, 16),
        (
            [c.nome, _reais(c.gastos), round(100 * c.gastos / total, 1), _reais(c.ganhos)]
            for c in d.categorias
        ),
        moeda=(2, 4),
    )
    _tabela(
        wb.create_sheet("Lançamentos"),
        ("Data", "Descrição", "Categoria", "Forma", "Tipo", "Valor", "Parcela"),
        (12, 34, 22, 18, 10, 14, 9),
        _linhas(d.lancamentos),
        moeda=(6,),
        datas=(1,),
    )
    _tabela(
        wb.create_sheet("Compromissos"),
        ("Vencimento", "Descrição", "Categoria", "Cartão/forma", "Tipo", "Valor", "Parcela"),
        (12, 34, 22, 18, 10, 14, 9),
        _linhas(d.compromissos),
        moeda=(6,),
        datas=(1,),
    )
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
